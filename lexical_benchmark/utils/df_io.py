from pathlib import Path

import pandas as pd
from openpyxl import Workbook, load_workbook


def df_write_xlsx(df_items: dict[str, pd.DataFrame], target: Path, sheet_name: str) -> None:
    """Create a XLSX file (or append into an existing one) with the given dataframes."""
    try:
        wb = load_workbook(str(target))
    except FileNotFoundError:
        wb = Workbook()

    ws = wb.create_sheet(title=sheet_name)
    row_counter = 0  # Row counter to track where to write the next dataframe
    for key, df in df_items.items():
        ws.cell(row=row_counter + 1, column=1).value = "-" * 5 + f"{key}" + "-" * 5
        ws.merge_cells(start_row=row_counter + 1, start_column=1, end_row=row_counter + 1, end_column=df.shape[1])
        row_counter += 1
        # Write column names
        ws.append(df.columns)
        row_counter += 1
        # Append rows of dataframe
        for row in df.itertuples(index=False, name=None):
            ws.append(row)
            row_counter += 1
    wb.save(str(target))
